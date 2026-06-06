"""
PowerPoint 解析器 — .pptx
"""

from pathlib import Path


class PPTParser:
    """解析 PowerPoint 文件，提取幻灯片文本（含表格）"""

    def parse(self, file_path: Path) -> list[dict]:
        from pptx import Presentation

        try:
            prs = Presentation(str(file_path))
        except Exception:
            return []

        results = []
        offset = 0
        for slide_num, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            texts.append(t)
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        line = " | ".join(c for c in cells if c)
                        if line:
                            texts.append(line)
            if texts:
                slide_text = f"[Slide {slide_num + 1}]\n" + "\n".join(texts)
                results.append({"content": slide_text.strip(), "char_offset": offset})
                offset += len(slide_text)
        return results
