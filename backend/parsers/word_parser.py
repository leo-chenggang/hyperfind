"""
Word 解析器 — .docx
"""

from pathlib import Path


class WordParser:
    """解析 Word 文档，提取段落和表格文本"""

    def parse(self, file_path: Path) -> list[dict]:
        from docx import Document

        try:
            doc = Document(str(file_path))
        except Exception:
            return []

        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    paragraphs.append(line)

        if not paragraphs:
            return []

        full_text = "\n".join(paragraphs)
        return [{"content": full_text, "char_offset": 0}]
